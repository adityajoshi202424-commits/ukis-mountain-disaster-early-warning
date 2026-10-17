import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
import os

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

s1 = prs.slides.add_slide(blank_layout)
s2 = prs.slides.add_slide(blank_layout)

morph_xml = (
    '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:p16="http://schemas.microsoft.com/office/powerpoint/2015/09/main" spd="med">'
    '<p16:morph option="byObject"/>'
    '</p:transition>'
)
transition_elem = parse_xml(morph_xml)
s2._element.append(transition_elem)

test_path = r"c:\Users\USER\Desktop\Adii Folder\sih ps flashflood\presentation\test_morph.pptx"
prs.save(test_path)
print("Saved test_morph.pptx successfully!")
